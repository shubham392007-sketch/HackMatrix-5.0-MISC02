import io
import pytest
from fastapi.testclient import TestClient
import fitz
import docx
import pptx

from backend.app.main import app
from backend.app.extraction.document_extractor import document_extractor

client = TestClient(app)
AUTH_HEADERS = {"Authorization": "Bearer sp_demo_alexchen_valid"}


def create_synthetic_pdf() -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 72), "SKILL PASSPORT PROFESSIONAL CERTIFICATE\n\nRecipient: Alex Chen\nCredential: Senior Distributed Systems Architect\nIssued By: Cloud Native Architecture Board\nDate: October 2026\nVerification Code: CERT-8899-CLOUD")
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def create_synthetic_docx() -> bytes:
    doc = docx.Document()
    doc.add_heading("Production Incident Postmortem", level=1)
    doc.add_paragraph("Summary: Mitigated database connection pool starvation during peak traffic.")
    doc.add_paragraph("Action Taken: Reconfigured keepalives and expanded connection pool.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def create_synthetic_pptx() -> bytes:
    prs = pptx.Presentation()
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "Microservices Architecture Migration"
    subtitle.text = "Q3 Engineering Retrospective by Alex Chen"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_extractor_pdf():
    pdf_bytes = create_synthetic_pdf()
    res = document_extractor.extract(pdf_bytes, "certificate.pdf", "application/pdf")
    assert res["page_count"] == 1
    assert "Alex Chen" in res["full_text"]
    assert "CERT-8899-CLOUD" in res["full_text"]
    assert len(res["chunks"]) == 1


def test_extractor_docx():
    docx_bytes = create_synthetic_docx()
    res = document_extractor.extract(docx_bytes, "postmortem.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert "Incident Postmortem" in res["full_text"] or "Incident Postmortem" in res["title"]
    assert "Mitigated database connection pool" in res["full_text"]


def test_extractor_pptx():
    pptx_bytes = create_synthetic_pptx()
    res = document_extractor.extract(pptx_bytes, "migration.pptx", "application/vnd.openxmlformats-officedocument.presentationml.presentation")
    assert "Microservices Architecture" in res["full_text"]
    assert res["page_count"] == 1


def test_upload_document_endpoint():
    pdf_bytes = create_synthetic_pdf()
    
    # Upload first time
    response = client.post(
        "/api/v1/documents/upload",
        files={"file": ("cert_test.pdf", pdf_bytes, "application/pdf")},
        data={"category": "certificate", "title": "Cloud Architect Certificate"},
        headers=AUTH_HEADERS,
    )
    assert response.status_code == 200
    data = response.json()
    assert data["status"] in ["success", "already_exists"]
    doc_id = data["document_id"]

    # Test deduplication
    dup_response = client.post(
        "/api/v1/documents/upload",
        files={"file": ("cert_test.pdf", pdf_bytes, "application/pdf")},
        data={"category": "certificate"},
        headers=AUTH_HEADERS,
    )
    assert dup_response.status_code == 200
    assert dup_response.json()["status"] == "already_exists"

    # List documents
    list_resp = client.get("/api/v1/documents", headers=AUTH_HEADERS)
    assert list_resp.status_code == 200
    assert list_resp.json()["total"] >= 1

    # Get single document
    get_resp = client.get(f"/api/v1/documents/{doc_id}", headers=AUTH_HEADERS)
    assert get_resp.status_code == 200
    assert get_resp.json()["id"] == doc_id
